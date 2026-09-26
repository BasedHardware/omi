import json
import tempfile
import unittest
from pathlib import Path
import sys

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "examples"))

from goals_to_xlsx import (
    convert,
    parse_time,
    sanitize_text,
    to_float,
)


class TestGoalsToXlsx(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dir_path = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_sanitize_text(self):
        self.assertEqual(sanitize_text("Clean Title"), "Clean Title")
        self.assertEqual(sanitize_text("=cmd|' /C calc'!A0"), "'=cmd|' /C calc'!A0")
        self.assertEqual(sanitize_text("@office"), "'@office")

    def test_to_float(self):
        self.assertEqual(to_float(10), 10.0)
        self.assertEqual(to_float("25.5"), 25.5)
        self.assertIsNone(to_float(None))
        self.assertIsNone(to_float("invalid"))

    def test_parse_time(self):
        dt = parse_time("2026-09-20T10:30:00Z")
        self.assertIsNotNone(dt)
        self.assertIsNone(dt.tzinfo)  # Must be naive for Excel
        self.assertEqual(dt.year, 2026)

    def test_convert_workbook_structure(self):
        goals = [
            {
                "id": "g1",
                "title": "Drink Water",
                "goal_type": "numeric",
                "current_value": 1.5,
                "target_value": 2.0,
                "unit": "liters",
                "is_active": True,
                "created_at": "2026-09-20T08:00:00Z",
                "updated_at": "2026-09-20T12:00:00Z",
            }
        ]
        in_file = self.dir_path / "goals.json"
        in_file.write_text(json.dumps(goals), encoding="utf-8")
        out_file = self.dir_path / "goals.xlsx"

        convert([str(in_file)], out_file)
        self.assertTrue(out_file.exists())

        wb = load_workbook(out_file)
        ws = wb.active
        self.assertEqual(ws.title, "Goals")
        self.assertEqual(ws.cell(row=1, column=1).value, "ID")
        self.assertEqual(ws.cell(row=2, column=1).value, "g1")
        self.assertEqual(ws.cell(row=2, column=2).value, "Drink Water")
        self.assertEqual(ws.cell(row=2, column=4).value, 1.5)
        self.assertEqual(ws.cell(row=2, column=5).value, 2.0)
        self.assertEqual(ws.cell(row=2, column=7).value, "=D2/E2")
        self.assertEqual(ws.cell(row=2, column=8).value, "Active")

        # Check overwrite refusal
        with self.assertRaises(FileExistsError):
            convert([str(in_file)], out_file)


if __name__ == "__main__":
    unittest.main()
