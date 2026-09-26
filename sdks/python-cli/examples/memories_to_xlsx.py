#!/usr/bin/env python3
"""
Convert Omi memories JSON export to an Excel workbook (.xlsx).

Usage:
    python memories_to_xlsx.py memories.json memories.xlsx
"""

import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

FIELDS = ("id", "category", "content", "tags", "visibility", "created_at")
DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"


def cell_text(value):
    """Render one exported field as text.

    Cells are written with explicit string type so values such as
    '=SUM(A1)' stay text and are never evaluated as formulas.
    """
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def cell_datetime(value):
    """Parse an ISO-8601 timestamp into a naive UTC datetime for Excel.

    Excel cells cannot carry a timezone, so every offset is converted to UTC
    and the header says so. Anything that is not a parseable timestamp is
    kept as text instead of being dropped.
    """
    if value is None:
        return None
    if not isinstance(value, str):
        return cell_text(value)
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return value
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
    return parsed


def convert(source, destination):
    raw_content = Path(source).read_text(encoding="utf-8")
    # Strip potential BOM
    if raw_content.startswith("\ufeff"):
        raw_content = raw_content[1:]
    items = json.loads(raw_content)

    if isinstance(items, dict):
        if "memories" in items and isinstance(items["memories"], list):
            items = items["memories"]
        elif "data" in items and isinstance(items["data"], list):
            items = items["data"]
        else:
            items = [items]

    if not isinstance(items, list):
        raise ValueError("Expected a JSON array or object from omi --json memory list")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "memories"
    header = [f"{name} (UTC)" if name.endswith("_at") else name for name in FIELDS]
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for item in items:
        if not isinstance(item, dict):
            raise ValueError("Each memory must be an object")

        created = cell_datetime(item.get("created_at"))
        tags_raw = item.get("tags") or []
        if isinstance(tags_raw, list):
            tags_str = ", ".join(str(t).strip() for t in tags_raw if t)
        else:
            tags_str = str(tags_raw)

        row = (
            cell_text(item.get("id")),
            cell_text(item.get("category") or "uncategorized"),
            cell_text(item.get("content")),
            cell_text(tags_str),
            cell_text(item.get("visibility") or "private"),
            created,
        )
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, datetime):
                cell.number_format = DATETIME_FORMAT
            elif isinstance(cell.value, str):
                cell.data_type = "s"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, name in enumerate(FIELDS, start=1):
        col_letter = get_column_letter(index)
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in sheet[col_letter])
        sheet.column_dimensions[col_letter].width = min(max(len(name), longest) + 2, 60)

    output_path = Path(destination)
    if output_path.exists():
        raise FileExistsError(f"Refusing to overwrite existing {output_path}")

    # Write next to the destination and rename atomically
    partial = output_path.with_name(output_path.name + ".partial")
    try:
        workbook.save(partial)
        os.replace(partial, output_path)
    except OSError:
        partial.unlink(missing_ok=True)
        raise


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Usage: python memories_to_xlsx.py INPUT.json OUTPUT.xlsx")
    try:
        convert(sys.argv[1], sys.argv[2])
    except (OSError, ValueError) as exc:
        sys.exit(f"XLSX export failed: {exc}")
