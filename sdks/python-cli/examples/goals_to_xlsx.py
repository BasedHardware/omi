#!/usr/bin/env python3
"""Convert Omi goal JSON exports into an Excel workbook (.xlsx)."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import sys
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = (
    "ID",
    "Title",
    "Goal Type",
    "Current Value",
    "Target Value",
    "Unit",
    "Progress %",
    "Status",
    "Min Value",
    "Max Value",
    "Created At (UTC)",
    "Updated At (UTC)",
)


def parse_time(value: Any) -> datetime | None:
    """Parse ISO-8601 timestamp string into naive UTC datetime for Excel."""
    if not isinstance(value, str) or not value.strip():
        return None
    normalized = value.strip().replace("Z", "+00:00")
    try:
        dt = datetime.fromisoformat(normalized)
    except ValueError:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    utc_dt = dt.astimezone(timezone.utc)
    return utc_dt.replace(tzinfo=None)  # Excel stores naive datetimes


def to_bool(value: Any) -> bool:
    """Normalize boolean or string flag."""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    if isinstance(value, str):
        return value.strip().lower() in ("true", "1", "yes", "active")
    return False


def to_float(value: Any) -> float | None:
    """Safely coerce to float."""
    if value is None:
        return None
    try:
        val = float(value)
        return val if sys.float_info.min <= abs(val) <= sys.float_info.max or val == 0.0 else None
    except (ValueError, TypeError, OverflowError):
        return None


def sanitize_text(value: Any) -> str:
    """Sanitize string values and defend against formula injection."""
    if value is None:
        return ""
    if not isinstance(value, str):
        value = json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value)
    clean = " ".join(value.split())
    if clean.startswith(("=", "+", "-", "@")):
        return "'" + clean
    return clean


def load(sources: Iterable[str | Path]) -> list[dict[str, Any]]:
    """Load and deduplicate goals across multiple export files."""
    items_by_id: dict[str, dict[str, Any]] = {}
    for source in sources:
        path = Path(source)
        content = path.read_bytes().decode("utf-8-sig")
        data = json.loads(content)
        raw_items = data.get("goals") or data.get("items") or data.get("data") or [data] if isinstance(data, dict) else data
        if not isinstance(raw_items, list):
            raise ValueError(f"{source}: expected JSON array or wrapped object containing goals")
        for item in raw_items:
            if not isinstance(item, dict):
                raise ValueError(f"{source}: each goal must be a JSON object")
            item_id = item.get("id")
            if not isinstance(item_id, str) or not item_id.strip():
                raise ValueError(f"{source}: goal missing valid string id")
            clean_id = item_id.strip()
            existing = items_by_id.get(clean_id)
            if existing is not None:
                new_dt = parse_time(item.get("updated_at") or item.get("created_at"))
                old_dt = parse_time(existing.get("updated_at") or existing.get("created_at"))
                if new_dt and old_dt:
                    if new_dt > old_dt:
                        items_by_id[clean_id] = item
                elif new_dt and not old_dt:
                    items_by_id[clean_id] = item
            else:
                items_by_id[clean_id] = item
    return list(items_by_id.values())


def convert(sources: list[str], output_path: str | Path) -> None:
    """Build workbook and write atomically to output path."""
    dest = Path(output_path)
    if dest.exists():
        raise FileExistsError(f"Refusing to overwrite existing {dest}")

    goals = load(sources)

    wb = Workbook()
    ws = wb.active
    ws.title = "Goals"

    # Header styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")

    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for row_idx, g in enumerate(goals, start=2):
        gid = sanitize_text(g.get("id"))
        title = sanitize_text(g.get("title"))
        gtype = sanitize_text(g.get("goal_type") or "qualitative")
        cur = to_float(g.get("current_value"))
        tgt = to_float(g.get("target_value"))
        unit = sanitize_text(g.get("unit"))
        is_act = to_bool(g.get("is_active", True))
        status = "Active" if is_act else "Inactive"
        min_v = to_float(g.get("min_value"))
        max_v = to_float(g.get("max_value"))
        created = parse_time(g.get("created_at"))
        updated = parse_time(g.get("updated_at"))

        ws.cell(row=row_idx, column=1, value=gid)
        ws.cell(row=row_idx, column=2, value=title)
        ws.cell(row=row_idx, column=3, value=gtype)
        ws.cell(row=row_idx, column=4, value=cur)
        ws.cell(row=row_idx, column=5, value=tgt)
        ws.cell(row=row_idx, column=6, value=unit)

        # Progress formula
        progress_cell = ws.cell(row=row_idx, column=7)
        if tgt and tgt > 0 and cur is not None:
            progress_cell.value = f"=D{row_idx}/E{row_idx}"
            progress_cell.number_format = "0.0%"
        else:
            progress_cell.value = None

        ws.cell(row=row_idx, column=8, value=status)
        ws.cell(row=row_idx, column=9, value=min_v)
        ws.cell(row=row_idx, column=10, value=max_v)

        cell_created = ws.cell(row=row_idx, column=11, value=created)
        if created:
            cell_created.number_format = "YYYY-MM-DD HH:MM:SS"

        cell_updated = ws.cell(row=row_idx, column=12, value=updated)
        if updated:
            cell_updated.number_format = "YYYY-MM-DD HH:MM:SS"

    # Freeze header and enable autofilter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = min(40, max(max_len + 3, 10))

    # Safe atomic write
    partial_path = dest.with_suffix(dest.suffix + ".partial")
    try:
        wb.save(partial_path)
        partial_path.replace(dest)
    except Exception:
        partial_path.unlink(missing_ok=True)
        raise


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", help="Destination .xlsx file path")
    parser.add_argument("inputs", nargs="+", help="Input goal JSON export files")
    args = parser.parse_args(argv)

    try:
        convert(args.inputs, args.output)
        print(f"Goals Excel workbook written to {args.output}")
        return 0
    except Exception as exc:
        sys.exit(f"Failed to generate Excel workbook: {exc}")


if __name__ == "__main__":
    raise SystemExit(main())
