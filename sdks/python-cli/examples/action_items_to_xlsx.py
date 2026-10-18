#!/usr/bin/env python3
"""Convert Omi action-item JSON exports to a real Excel workbook (.xlsx).

See action_items_xlsx.md for the full recipe. Requires ``openpyxl``
(``pip install openpyxl``) — the only non-stdlib piece, matching the
conversations XLSX recipe.

Usage:
    omi --json action-item list --limit 500 > action_items.json
    python action_items_to_xlsx.py action_items.json [page2.json ...] -o action_items.xlsx

due_at/created_at/updated_at land as native Excel datetime cells (UTC),
``completed`` as a native Excel boolean, and every string is explicitly
string-typed so a description like ``=SUM(A1)`` is never executed as a
formula. AutoFilter + frozen header are on; columns auto-size.
"""

import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

FIELDS = ("id", "description", "completed", "due_at", "created_at", "updated_at", "conversation_id", "source")
DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"
ENVELOPE_KEYS = ("action_items", "items", "data")


def cell_text(value: Any) -> str:
    """Render one exported field as spreadsheet-safe text.

    The dev API is loosely typed, so a field can arrive as a non-string;
    anything non-null is coerced rather than rejected so one odd row cannot
    destroy the whole export.
    """
    if value is None:
        return ""
    if not isinstance(value, str):
        value = json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value)
    return " ".join(value.split())


def cell_datetime(value: Any):
    """Parse an ISO-8601 timestamp into a naive UTC datetime for Excel.

    Excel cells cannot carry a timezone, so every offset is converted to UTC
    and the header says so. Anything that is not a parseable timestamp is
    kept as text instead of being dropped.
    """
    if value is None:
        return None
    if not isinstance(value, str):
        return cell_text(value)
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return value
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
    return parsed


def cell_bool(value: Any) -> Optional[bool]:
    """Loose boolean → real Excel TRUE/FALSE; non-boolean → None (blank)."""
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        low = value.strip().lower()
        if low in ("true", "yes", "1", "t"):
            return True
        if low in ("false", "no", "0", "f"):
            return False
    if isinstance(value, (int, float)):
        return bool(value)
    return None


def extract_items(data: Any) -> List[Any]:
    """Unwrap items from bare lists or wrapped envelope dictionaries."""
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        for key in ENVELOPE_KEYS:
            if isinstance(data.get(key), list):
                return data[key]
        return [data]
    return []


def load(sources: List[str]) -> List[Dict[str, Any]]:
    """Read every page; dedupe by action-item id (first occurrence wins)."""
    seen = {}
    for src in sources:
        raw = json.loads(Path(src).read_bytes().decode("utf-8-sig"))
        for item in extract_items(raw):
            if not isinstance(item, dict):
                continue
            iid = str(item.get("id") or item.get("action_item_id") or "")
            if iid and iid not in seen:
                seen[iid] = item
            elif not iid:
                seen.setdefault(f"__noid_{len(seen)}", item)
    return list(seen.values())


def convert(sources: List[str], destination: str) -> int:
    """Write the workbook; return the row count."""
    items = load(sources)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "action_items"
    header = [f"{name} (UTC)" if name.endswith("_at") else name for name in FIELDS]
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for item in items:
        row = (
            cell_text(item.get("id") or item.get("action_item_id")),
            cell_text(item.get("description") or item.get("title")),
            cell_bool(item.get("completed")),
            cell_datetime(item.get("due_at")),
            cell_datetime(item.get("created_at")),
            cell_datetime(item.get("updated_at")),
            cell_text(item.get("conversation_id")),
            cell_text(item.get("source") or "omi"),
        )
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, datetime):
                cell.number_format = DATETIME_FORMAT
            elif isinstance(cell.value, str):
                cell.data_type = "s"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, name in enumerate(FIELDS, start=1):
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in sheet[get_column_letter(index)])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(name), longest) + 2, 60)

    output_path = Path(destination)
    if ".." in output_path.parts:
        raise ValueError(f"Refusing '..' in output path {destination}")
    if output_path.exists():
        raise FileExistsError(f"Refusing to overwrite existing {output_path}")
    # Write next to the destination and rename, so a failed save cannot leave a
    # truncated workbook behind for the next run.
    partial = output_path.with_name(output_path.name + ".partial")
    try:
        workbook.save(partial)
        os.replace(partial, output_path)
    except Exception:
        partial.unlink(missing_ok=True)
        raise
    return len(items)


def main(argv: Optional[List[str]] = None) -> int:
    import argparse

    ap = argparse.ArgumentParser(description="Convert Omi action-item JSON exports to .xlsx.")
    ap.add_argument("inputs", nargs="+", help="One or more JSON export files")
    ap.add_argument("-o", "--output", required=True, help="Destination .xlsx path")
    args = ap.parse_args(argv)
    try:
        n = convert(args.inputs, args.output)
    except (ValueError, FileExistsError, OSError, json.JSONDecodeError) as e:
        print(f"action_items_to_xlsx: {e}", file=sys.stderr)
        return 2
    print(f"action_items_to_xlsx: wrote {n} row(s) to {args.output}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
